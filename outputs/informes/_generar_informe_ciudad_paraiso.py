# -*- coding: utf-8 -*-
"""
Genera el informe Word del proceso ITT para Ciudad Paraiso (Comuna 3),
a partir del notebook notebooks/08_itt_ciudad_paraiso_5dim.ipynb.

Salida: outputs/informes/Informe_ITT_Ciudad_Paraiso.docx
"""
from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- paleta
AZUL   = RGBColor(0x1B, 0x4F, 0x8A)   # azul institucional (header del Excel del notebook)
VERDE  = RGBColor(0x00, 0x9E, 0x73)   # Okabe-Ito verde
NARANJA = RGBColor(0xE6, 0x9F, 0x00)  # Okabe-Ito naranja
BERMELL = RGBColor(0xD5, 0x5E, 0x00)  # Okabe-Ito bermellon
GRIS   = RGBColor(0x33, 0x33, 0x33)
GRIS_T = RGBColor(0x6C, 0x75, 0x7D)

OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "Informe_ITT_Ciudad_Paraiso.docx"


# ---------------------------------------------------------------- helpers
def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def style_header_row(row, fill="1B4F8A"):
    for cell in row.cells:
        set_cell_bg(cell, fill)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.bold = True
                r.font.size = Pt(9)


def add_table(doc, headers, rows, widths=None, header_fill="1B4F8A",
              body_size=9, zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
    style_header_row(table.rows[0], header_fill)

    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            run = p.add_run("" if val is None else str(val))
            run.font.size = Pt(body_size)
            if zebra and ri % 2 == 1:
                set_cell_bg(cells[ci], "F4F6F9")

    if widths:
        for row in table.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Inches(w)
    return table


def h_title(doc, text, color=AZUL, size=22):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    return p


def h1(doc, text):
    p = doc.add_heading(level=1)
    run = p.add_run(text)
    run.font.color.rgb = AZUL
    run.font.size = Pt(15)
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    run.font.color.rgb = GRIS
    run.font.size = Pt(12)
    return p


def para(doc, text, size=10.5, italic=False, color=None, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return p


def bullet(doc, text, size=10.5, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.bold = True
        r.font.size = Pt(size)
    r2 = p.add_run(text)
    r2.font.size = Pt(size)
    return p


# ================================================================ documento
doc = Document()

# estilo base
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

# ---------------------------------------------------------------- portada
h_title(doc, "Informe del Proceso ITT")
p = doc.add_paragraph()
r = p.add_run("Ciudad Paraíso — Comuna 3, Cali")
r.font.size = Pt(15)
r.font.color.rgb = VERDE
r.font.bold = True

para(doc, "Índice de Transformación Territorial (ITT) · Modelo de 5 dimensiones",
     size=11, color=GRIS_T)
para(doc, "Fuente: notebooks/08_itt_ciudad_paraiso_5dim.ipynb", size=10, italic=True, color=GRIS_T)
para(doc, f"Fecha de elaboración: {date.today().strftime('%d/%m/%Y')}", size=10, italic=True, color=GRIS_T)
para(doc, "Período analizado: 2023 – 2026 (2026 parcial, no comparable)", size=10, italic=True, color=GRIS_T)

doc.add_paragraph()

# ---------------------------------------------------------------- 1. Resumen
h1(doc, "1. Resumen ejecutivo")
para(doc,
     "Este informe documenta el proceso de cálculo del Índice de Transformación Territorial (ITT) "
     "para la zona de renovación urbana Ciudad Paraíso (Comuna 3 de Cali, que agrupa los barrios "
     "San Pascual, El Calvario, San Juan Bosco y Santa Rosa). El notebook replica la estructura y "
     "metodología de la referencia vigente del repositorio (04_itt_barrio_obrero_5dim.ipynb) e "
     "incorpora dos mejoras explícitas: (1) verificación espacial real por point-in-polygon antes de "
     "agregar cualquier indicador, y (2) cálculo de Desarrollo Económico con umbrales fijos "
     "ref_min/ref_max en vez de min-max relativo de la propia muestra.")
para(doc,
     "El ITT opera en escala 0–100: a mayor valor, mayor transformación territorial. Se calcula como "
     "suma ponderada de 5 dimensiones, cada una normalizada a 0–100 con umbrales fijos y promediada "
     "a partir de sus indicadores.")

# ---------------------------------------------------------------- 2. Alcance
h1(doc, "2. Alcance y supuestos declarados")
bullet(doc, "Educación y Desarrollo NO se calcula en este repositorio; se gestiona en otro repositorio "
            "(indicación del usuario, 2026-10-04). La dimensión DesSocial de este notebook se calcula solo "
            "con conflictividad (VIF + Riñas + SPA) vía nanmean, sin alterar la fórmula ni los pesos del ITT.",
       bold_prefix="Educación fuera de alcance: ")
bullet(doc, "La zona no tiene NDVI ni censo arbóreo propio (no existe carpeta 5_Entorno_Urbano en los datos). "
            "Se usa un proxy de déficit habitacional de la Comuna 3 (corte único 2024, sin serie temporal).",
       bold_prefix="Entorno Urbano por proxy: ")
bullet(doc, "El año 2026 tiene cobertura incompleta en todas las fuentes (DATIC hasta Q1, siniestros hasta Q2). "
            "Su ITT se calcula y se muestra como informativo, pero se marca «Parcial (no comparable)» y no se le "
            "asigna un nivel real de transformación.",
       bold_prefix="2026 parcial: ")

# ---------------------------------------------------------------- 3. Dimensiones y pesos
h1(doc, "3. Dimensiones, pesos e indicadores")
para(doc, "Los pesos de las 5 dimensiones suman 1.00. En esta zona se mantienen idénticos a Barrio Obrero "
          "para preservar comparabilidad metodológica.")
add_table(
    doc,
    ["Dimensión", "Peso", "Indicadores usados en esta zona", "Tipo"],
    [
        ["Seguridad", "27%", "Homicidios, Hurtos", "Inverso"],
        ["Movilidad", "22%", "Siniestralidad, Lesionados, Mortales", "Inverso"],
        ["DesSocial (Cohesión)", "19%", "VIF, Riñas, SPA (Educación fuera de alcance)", "Inverso"],
        ["Entorno Urbano", "17%", "Proxy déficit habitacional Comuna 3 (corte 2024)", "Proxy fijo"],
        ["Desarrollo Económico", "15%", "Establecimientos, Empleabilidad, Ingresos (log1p)", "Positivo"],
    ],
    widths=[1.7, 0.6, 3.3, 0.9],
)

# ---------------------------------------------------------------- 4. Umbrales
h1(doc, "4. Umbrales de normalización (ref_min / ref_max)")
para(doc, "Calibrados sobre el rango real observado 2023-2025 en Ciudad Paraíso (no reutilizados de Barrio "
          "Obrero, cuyos umbrales saturarían esta zona en score 0). Todos son inversos: menor valor = mejor "
          "score. Pendientes de validación por un experto local del territorio.")
add_table(
    doc,
    ["Indicador", "ref_min", "ref_max", "Rango observado 2023-2025"],
    [
        ["Homicidios", "0", "20", "6 – 16 / año"],
        ["Hurtos", "80", "320", "164 – 260 / año"],
        ["Siniestralidad", "10", "45", "24 – 35 / año"],
        ["Lesionados", "10", "38", "23 – 30 / año"],
        ["Mortales", "0", "6", "0 – 4 / año"],
        ["VIF", "5", "25", "11 – 17 / año"],
        ["Riñas", "150", "500", "327 – 418 / año"],
        ["SPA", "10", "150", "28 – 121 / año"],
    ],
    widths=[1.8, 1.1, 1.1, 2.5],
)
para(doc, "Desarrollo Económico (umbrales fijos, positivos): establecimientos [100–450], "
          "empleabilidad [300–1200], log1p(ingresos) [24–28].", size=9.5, italic=True, color=GRIS_T)
para(doc, "Fórmula de normalización: score_raw = clamp((valor − ref_min) / (ref_max − ref_min) × 100, 0, 100); "
          "para indicadores inversos, score = 100 − score_raw.", size=9.5, italic=True, color=GRIS_T)

# ---------------------------------------------------------------- 5. Proceso
h1(doc, "5. Proceso de cálculo paso a paso")
para(doc, "El notebook ejecuta el flujo en el siguiente orden (celdas reales del notebook):")

add_table(
    doc,
    ["Celda", "Etapa", "Qué hace"],
    [
        ["1 / 1B", "Preparación", "Instala dependencias y clona/actualiza el repositorio en Colab."],
        ["2", "Configuración", "Importaciones y paleta visual Okabe-Ito."],
        ["3", "Parámetros", "Rutas, pesos, umbrales ref_min/ref_max, años, comuna proxy."],
        ["3B", "Entorno Urbano", "Calcula el proxy de déficit habitacional (Comuna 3, corte 2024)."],
        ["4", "Carga", "Lee GeoJSON de las 6 fuentes + polígono de la zona."],
        ["4B", "Verificación espacial", "point-in-polygon real con geopandas.sjoin (no confía en nom_barrio)."],
        ["5", "Mapa interactivo", "Mapa multicapa Folium (capas base, red peatonal OSM, eventos)."],
        ["6", "Procesamiento", "Agrega indicadores anual y trimestral; enmascara 2026 sin dato (NaN)."],
        ["7", "Normalización", "Scores por indicador y por dimensión (Seguridad, Movilidad, DesSocial)."],
        ["DE / DE-Trim", "Desarrollo Económico", "Score por stock acumulado con umbrales fijos + heatmap trimestral."],
        ["15–18", "Consolidación", "ITT global, heatmaps, evolución trimestral, radar y composición."],
        ["19", "Exportación", "Exporta a Excel (9 hojas) en outputs/itt_ciudad_paraiso."],
    ],
    widths=[0.9, 1.6, 4.0],
)

h2(doc, "5.1 Agregación del ITT")
para(doc, "Paso 1 — cada indicador se normaliza a 0–100. Paso 2 — cada dimensión promedia sus indicadores "
          "(nanmean, que ignora componentes sin dato). Paso 3 — el ITT es la suma ponderada:")
para(doc, "ITT = 0.27·Seguridad + 0.22·Movilidad + 0.19·DesSocial + 0.17·EntornoU + 0.15·DesEco",
     size=10, bold=True, color=AZUL)

h2(doc, "5.2 Clasificación por nivel")
add_table(
    doc,
    ["Rango ITT", "Nivel"],
    [
        ["0 – 40", "Activación"],
        ["40 – 60", "Consolidación"],
        ["60 – 80", "Transformación"],
        ["80 – 100", "Escala"],
    ],
    widths=[1.5, 2.5],
)

# ---------------------------------------------------------------- 6. Verificación espacial
h1(doc, "6. Verificación espacial (control de calidad)")
para(doc, "A diferencia de otras zonas, en Ciudad Paraíso los campos nom_comuna / BARRIO de las fuentes son "
          "poco confiables (ej. comparendos etiquetados como «C-22 LA MARIA» para puntos físicamente en El "
          "Calvario; Registro Mercantil con negocios en 14 comunas distintas). La Celda 4B hace un cruce "
          "espacial real contra la geometría del polígono y usa ese resultado como único filtro de verdad. "
          "Si una fuente no llega a ~100% dentro del polígono, debe filtrarse por su máscara antes de agregar.")

# ---------------------------------------------------------------- 7. Resultados
h1(doc, "7. Resultados")
para(doc, "Los valores numéricos del ITT dependen de ejecutar el notebook con los datos del repositorio. "
          "Este informe documenta el proceso y la estructura de salida; no reproduce cifras que no fueron "
          "ejecutadas aquí. Al correr el notebook, los resultados quedan consolidados en:")
bullet(doc, "outputs/itt_ciudad_paraiso/ITT_Ciudad_Paraiso_Consolidado.xlsx — 9 hojas "
            "(ITT_Resumen, Seguridad, Movilidad, Cohesion_Social_DesSocial, Entorno_Urbano_proxy, "
            "Educacion_fuera_alcance, Desarrollo_Economico, Indicadores_Trimestrales, Auditoria_Spatial_Join).")
bullet(doc, "Figuras PNG (heatmaps por dimensión, evolución trimestral línea+área, radar de 5 dimensiones, "
            "ITT global y composición, dashboard del proxy de Entorno Urbano).")

doc.add_page_break()

# ---------------------------------------------------------------- 8. Tabla de seguimiento
h1(doc, "8. Tabla de seguimiento — ¿qué falta?")
para(doc, "Estado de cada componente del ITT para esta zona. Leyenda: Completo = listo · Parcial = en progreso o "
          "provisional · Pendiente / Fuera de alcance = sin implementar en este repo.")

add_table(
    doc,
    ["Componente", "Estado", "Detalle / qué falta"],
    [
        ["Polígono de la zona", "Completo", "Cargado y reproyectado a WGS84. Área real calculada en EPSG:3115."],
        ["Verificación espacial (4B)", "Completo", "sjoin point-in-polygon implementado para las 6 fuentes."],
        ["Seguridad", "Completo", "Homicidios + Hurtos con umbrales propios 2023-2025."],
        ["Movilidad", "Completo", "Siniestralidad + Lesionados + Mortales (siniestros cubren hasta 2026-Q2)."],
        ["DesSocial / Cohesión", "Parcial", "VIF + Riñas + SPA calculados. Falta el componente de Educación "
                                            "(fuera de alcance de este repo)."],
        ["Entorno Urbano", "Parcial", "Solo proxy de déficit habitacional (Comuna 3, corte único 2024). "
                                      "Falta NDVI / área verde / censo arbóreo propio de la zona."],
        ["Desarrollo Económico", "Parcial", "Calculado con stock acumulado y umbrales fijos. 2026 usa 2025 "
                                            "como referencia (no hay corte 2026). Umbrales pendientes de validación."],
        ["Educación y Desarrollo", "Fuera de alcance", "No se calcula en este repositorio; se gestiona en otro repo."],
        ["Cobertura 2026", "Parcial", "DATIC solo Q1; siniestros hasta Q2. Año marcado «no comparable»."],
        ["Validación de umbrales", "Pendiente", "ref_min/ref_max requieren validación por experto local del territorio."],
        ["Mapa interactivo (Celda 5)", "Completo", "Multicapa Folium (5 bases, red peatonal OSM, nodos, heatmap, eventos)."],
        ["Exportación a Excel", "Completo", "9 hojas con formato en outputs/itt_ciudad_paraiso."],
        ["Datos VBG", "Pendiente", "Cargado (raw_vbg) pero aún no incorporado al cálculo de ninguna dimensión."],
    ],
    widths=[1.9, 1.1, 3.5],
)

h2(doc, "8.1 Prioridades sugeridas")
bullet(doc, "Validar los umbrales ref_min/ref_max con conocimiento del territorio (impacto directo en los scores).")
bullet(doc, "Conseguir datos propios de Entorno Urbano (NDVI / área verde) para reemplazar el proxy fijo.")
bullet(doc, "Definir si VBG entra al cálculo de Cohesión o se mantiene como contexto.")
bullet(doc, "Reemplazar la referencia 2025 de Desarrollo Económico cuando exista corte 2026.")

# ---------------------------------------------------------------- cierre
doc.add_paragraph()
para(doc, "Documento generado automáticamente a partir del notebook de Ciudad Paraíso. "
          "Las cifras definitivas deben leerse de la ejecución del notebook y del Excel consolidado.",
     size=9, italic=True, color=GRIS_T)

doc.save(str(DOCX_PATH))
print(f"Informe generado: {DOCX_PATH}")
